#!/usr/bin/env python3
"""
Tests for billing_export.py.

Covers:
- get_billing_data() follows NextPageToken across pages and sums a
  (month, service) pair split across pages instead of overwriting it
- COST_METRIC is the single key used for the request and the parse
- validate_date_input() returns an EXCLUSIVE end date (Cost Explorer
  TimePeriod.End is exclusive), so the final day of the final month is kept
- create_excel_report() Summary sheet contract (Month | Total Cost (USD),
  trailing 'Total All Months' row)
- AccessDenied path: exit 0, warning logged, skip marker written; the marker
  is picked up by the Smart Scan executor and lands in the zip, and never
  matches the real billing export glob

moto's Cost Explorer does not model pagination, so the CE client is a real
botocore client wrapped in botocore.stub.Stubber (validates request params
and response shapes against the service model).
"""

import datetime
import fnmatch
import importlib.util
import sys
import zipfile
from pathlib import Path
from unittest.mock import patch

import boto3
import pytest
from botocore.stub import Stubber
from openpyxl import load_workbook

ROOT = Path(__file__).parent.parent.parent
sys.path.insert(0, str(ROOT / "scripts"))
import billing_export  # noqa: E402

REAL_EXPORT_GLOB = "*billing-last-12-months-export-*.xlsx"


@pytest.fixture(autouse=True)
def fake_aws_credentials(monkeypatch):
    monkeypatch.setenv("AWS_ACCESS_KEY_ID", "testing")
    monkeypatch.setenv("AWS_SECRET_ACCESS_KEY", "testing")
    monkeypatch.setenv("AWS_SESSION_TOKEN", "testing")
    monkeypatch.setenv("AWS_DEFAULT_REGION", "us-east-1")


@pytest.fixture(autouse=True)
def patch_output_dir(tmp_path, monkeypatch):
    monkeypatch.setattr(billing_export.utils, "get_output_dir", lambda: tmp_path)
    yield tmp_path


def _stubbed_ce(monkeypatch):
    client = boto3.client("ce", region_name="us-east-1")
    stubber = Stubber(client)
    monkeypatch.setattr(billing_export.utils, "get_boto3_client", lambda *a, **k: client)
    return client, stubber


def _group(service, amount):
    return {
        "Keys": [service],
        "Metrics": {billing_export.COST_METRIC: {"Amount": amount, "Unit": "USD"}},
    }


def _result(start, end, groups):
    return {
        "TimePeriod": {"Start": start, "End": end},
        "Total": {},
        "Groups": groups,
        "Estimated": False,
    }


def _base_request(start, end):
    return {
        "TimePeriod": {"Start": start, "End": end},
        "Granularity": "MONTHLY",
        "Metrics": [billing_export.COST_METRIC],
        "GroupBy": [{"Type": "DIMENSION", "Key": "SERVICE"}],
    }


def _months_in(start, end):
    """'YYYY-MM' keys for every calendar month in [start, end)."""
    months = []
    y, m = start.year, start.month
    while datetime.datetime(y, m, 1) < end:
        months.append(f"{y:04d}-{m:02d}")
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)
    return months


def _frozen_now(fixed):
    """Patch billing_export's datetime.datetime.now() to a fixed instant."""

    class FakeDT(datetime.datetime):
        @classmethod
        def now(cls, tz=None):
            return fixed

    return patch.object(billing_export.datetime, "datetime", FakeDT)


class TestPagination:
    def test_follows_next_page_token_and_sums_split_groups(self, monkeypatch):
        _, stubber = _stubbed_ce(monkeypatch)
        start, end = "2026-07-01", "2026-09-01"

        # Page 1: all of July, first part of August (EC2 split across pages)
        stubber.add_response(
            "get_cost_and_usage",
            {
                "ResultsByTime": [
                    _result("2026-07-01", "2026-08-01", [_group("Amazon EC2", "10.00")]),
                    _result("2026-08-01", "2026-09-01", [
                        _group("Amazon EC2", "100.25"),
                        _group("Amazon S3", "5.00"),
                    ]),
                ],
                "NextPageToken": "page-2",
            },
            _base_request(start, end),
        )
        # Page 2: remainder of August, EC2 again
        stubber.add_response(
            "get_cost_and_usage",
            {
                "ResultsByTime": [
                    _result("2026-08-01", "2026-09-01", [
                        _group("Amazon EC2", "50.50"),
                        _group("AWS Lambda", "1.25"),
                    ]),
                ],
                "NextPageToken": "page-3",
            },
            {**_base_request(start, end), "NextPageToken": "page-2"},
        )
        # Page 3: last page, no token
        stubber.add_response(
            "get_cost_and_usage",
            {
                "ResultsByTime": [
                    _result("2026-08-01", "2026-09-01", [_group("Amazon S3", "0.75")]),
                ],
            },
            {**_base_request(start, end), "NextPageToken": "page-3"},
        )

        with stubber:
            data = billing_export.get_billing_data(
                datetime.datetime(2026, 7, 1), datetime.datetime(2026, 9, 1)
            )
            stubber.assert_no_pending_responses()

        assert data["2026-07"] == {"Amazon EC2": pytest.approx(10.00)}
        assert data["2026-08"]["Amazon EC2"] == pytest.approx(150.75)
        assert data["2026-08"]["Amazon S3"] == pytest.approx(5.75)
        assert data["2026-08"]["AWS Lambda"] == pytest.approx(1.25)
        assert sum(data["2026-08"].values()) == pytest.approx(157.75)

    def test_single_page_no_token(self, monkeypatch):
        _, stubber = _stubbed_ce(monkeypatch)
        stubber.add_response(
            "get_cost_and_usage",
            {"ResultsByTime": [
                _result("2026-08-01", "2026-09-01", [_group("Amazon EC2", "3.00")]),
            ]},
            _base_request("2026-08-01", "2026-09-01"),
        )
        with stubber:
            data = billing_export.get_billing_data(
                datetime.datetime(2026, 8, 1), datetime.datetime(2026, 9, 1)
            )
            stubber.assert_no_pending_responses()
        assert data == {"2026-08": {"Amazon EC2": pytest.approx(3.00)}}


class TestMetricConstant:
    def test_metric_value_unchanged(self):
        assert billing_export.COST_METRIC == "NetUnblendedCost"

    def test_metric_is_valid_for_operation(self):
        # Botocore's model documents the valid Metrics values for this op.
        client = boto3.client("ce", region_name="us-east-1")
        doc = client.meta.service_model.operation_model("GetCostAndUsage") \
            .input_shape.members["Metrics"].documentation
        assert f"<code>{billing_export.COST_METRIC}</code>" in doc

    def test_no_stray_metric_literals(self):
        source = (ROOT / "scripts" / "billing_export.py").read_text()
        # Only the constant definition may spell the metric out.
        for metric in ("NetUnblendedCost", "BlendedCost", "UnblendedCost"):
            quoted = source.count(f"'{metric}'") + source.count(f'"{metric}"')
            assert quoted == (1 if metric == billing_export.COST_METRIC else 0), metric


class TestEndDateExclusive:
    def test_single_month_end_is_first_of_next_month(self):
        ok, _, start, end = billing_export.validate_date_input("08-2026")
        assert ok
        assert start == datetime.datetime(2026, 8, 1)
        assert end == datetime.datetime(2026, 9, 1)

    def test_december_rolls_year(self):
        _, _, start, end = billing_export.validate_date_input("12-2025")
        assert start == datetime.datetime(2025, 12, 1)
        assert end == datetime.datetime(2026, 1, 1)

    def test_last_12_end_is_first_of_current_month(self):
        today = datetime.datetime.now()
        _, _, start, end = billing_export.validate_date_input("last 12")
        assert end == datetime.datetime(today.year, today.month, 1)
        assert start == datetime.datetime(today.year - 1, today.month, 1)
        assert len(_months_in(start, end)) == 12

    @pytest.mark.parametrize("now, exp_start, exp_end, first, last", [
        # mid-month
        (datetime.datetime(2026, 9, 15, 10, 0), datetime.datetime(2025, 9, 1),
         datetime.datetime(2026, 9, 1), "2025-09", "2026-08"),
        # 1st of the month, just after midnight
        (datetime.datetime(2026, 9, 1, 0, 5), datetime.datetime(2025, 9, 1),
         datetime.datetime(2026, 9, 1), "2025-09", "2026-08"),
        # January: window is the whole previous calendar year
        (datetime.datetime(2026, 1, 20, 9, 0), datetime.datetime(2025, 1, 1),
         datetime.datetime(2026, 1, 1), "2025-01", "2025-12"),
        # January 1st
        (datetime.datetime(2026, 1, 1, 0, 1), datetime.datetime(2025, 1, 1),
         datetime.datetime(2026, 1, 1), "2025-01", "2025-12"),
        # February: previous month is January of the same year
        (datetime.datetime(2026, 2, 1, 12, 0), datetime.datetime(2025, 2, 1),
         datetime.datetime(2026, 2, 1), "2025-02", "2026-01"),
    ])
    def test_last_12_is_exactly_12_months_and_passes_retention(
        self, now, exp_start, exp_end, first, last
    ):
        with _frozen_now(now):
            ok, _, start, end = billing_export.validate_date_input("last 12")
            valid, message, retention = billing_export.validate_date_range(start, end)
        assert ok
        assert (start, end) == (exp_start, exp_end)
        months = _months_in(start, end)
        assert len(months) == 12
        assert months[0] == first and months[-1] == last
        assert retention == 14
        assert valid, message

    def test_previous_month_valid_on_first_of_month(self):
        fake_now = datetime.datetime(2026, 9, 1, 0, 30)

        class FakeDT(datetime.datetime):
            @classmethod
            def now(cls, tz=None):
                return fake_now

        with patch.object(billing_export.datetime, "datetime", FakeDT):
            ok, message, _ = billing_export.validate_date_range(
                datetime.datetime(2026, 8, 1), datetime.datetime(2026, 9, 1)
            )
        assert ok, message

    def test_current_month_rejected(self):
        today = datetime.datetime.now()
        start = datetime.datetime(today.year, today.month, 1)
        end = (start + datetime.timedelta(days=32)).replace(day=1)
        ok, _, _ = billing_export.validate_date_range(start, end)
        assert not ok

    def test_request_uses_exclusive_end(self, monkeypatch):
        _, stubber = _stubbed_ce(monkeypatch)
        _, _, start, end = billing_export.validate_date_input("08-2026")
        stubber.add_response(
            "get_cost_and_usage",
            {"ResultsByTime": []},
            _base_request("2026-08-01", "2026-09-01"),
        )
        with stubber:
            billing_export.get_billing_data(start, end)
            stubber.assert_no_pending_responses()


class TestSummaryContract:
    DATA = {
        "2026-07": {"Amazon EC2": 10.0},
        "2026-08": {"Amazon EC2": 150.75, "Amazon S3": 5.75},
    }

    def _report(self):
        return billing_export.create_excel_report(
            self.DATA, "TEST-ACCOUNT", "last-12-months",
            account_id="123456789012",
            start_date=datetime.datetime(2026, 7, 1),
            end_date=datetime.datetime(2026, 9, 1),
        )

    def test_sheet_order(self, patch_output_dir):
        wb = load_workbook(self._report())
        assert wb.sheetnames == ["Summary", "About", "Savings Plans", "Jul 2026", "Aug 2026"]

    def test_about_sheet(self, patch_output_dir, monkeypatch):
        monkeypatch.setattr(billing_export.utils, "get_version", lambda: "9.9.9-test")
        wb = load_workbook(self._report())
        rows = list(wb["About"].iter_rows(values_only=True))
        assert rows[0] == ("Field", "Value")
        about = dict(rows[1:])
        assert [r[0] for r in rows[1:]] == [
            "Account ID", "Account Name", "Cost Metric", "Metric Meaning",
            "Period Start (inclusive)", "Period End (inclusive)", "Generated",
            "Tool Version", "Data Scope", "Org Caveat", "GovCloud Note",
            "Invoice Caveat", "API Cost Note",
        ]
        assert about["Account ID"] == "123456789012"
        assert about["Account Name"] == "TEST-ACCOUNT"
        assert about["Cost Metric"] == "NetUnblendedCost"
        assert about["Period Start (inclusive)"] == "2026-07-01"
        assert about["Period End (inclusive)"] == "2026-08-31"
        assert about["Tool Version"] == "9.9.9-test"
        for key in ("Metric Meaning", "Generated", "Data Scope", "Org Caveat", "GovCloud Note"):
            assert about[key]
        assert about["Invoice Caveat"] == billing_export.INVOICE_CAVEAT
        assert "reseller" in about["Invoice Caveat"]
        assert "cannot see that invoice" in about["Invoice Caveat"]
        assert "paginated Cost Explorer API request" in about["API Cost Note"]
        assert "RECORD_TYPE" in about["API Cost Note"]

    def test_month_sheets_unchanged(self, patch_output_dir):
        wb = load_workbook(self._report())
        rows = list(wb["Aug 2026"].iter_rows(values_only=True))
        assert rows[0] == ("Service", "Cost (USD)")
        assert rows[1] == ("Amazon EC2", pytest.approx(150.75))
        assert rows[2] == ("Amazon S3", pytest.approx(5.75))
        assert rows[-1] == ("Total", pytest.approx(156.5))

    def test_month_sheet_names_cannot_be_about(self):
        for year in (2000, 2026, 2099):
            for month in range(1, 13):
                name = datetime.datetime(year, month, 1).strftime("%b %Y")
                assert name.lower() != billing_export.ABOUT_SHEET.lower()

    def test_summary_sheet_shape(self, patch_output_dir):
        path = self._report()
        assert fnmatch.fnmatch(Path(path).name, REAL_EXPORT_GLOB)

        wb = load_workbook(path)
        ws = wb["Summary"]
        rows = list(ws.iter_rows(values_only=True))
        assert rows[0] == ("Month", "Total Cost (USD)")
        assert rows[1] == ("July 2026", pytest.approx(10.0))
        assert rows[2][0] == "August 2026"
        assert rows[2][1] == pytest.approx(156.5)
        assert rows[3][0] == "Total All Months"
        assert rows[3][1] == pytest.approx(166.5)
        assert len(rows) == 4


def _import_smart_scan_package():
    """Import scripts/smart_scan (package), not the root smart_scan.py CLI.

    Both share a name; whichever directory is first on sys.path wins, so pin
    scripts/ to the front and drop a shadowing non-package module.
    """
    scripts_dir = str(ROOT / "scripts")
    if sys.path[0] != scripts_dir:
        sys.path.insert(0, scripts_dir)
    cached = sys.modules.get("smart_scan")
    if cached is not None and not hasattr(cached, "__path__"):
        for name in [n for n in sys.modules if n == "smart_scan" or n.startswith("smart_scan.")]:
            sys.modules.pop(name, None)
    import smart_scan.executor as executor_mod
    return executor_mod


def _load_smart_scan_module():
    _import_smart_scan_package()
    spec = importlib.util.spec_from_file_location("_smart_scan_cli_for_test", ROOT / "smart_scan.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


class TestPermissionSkip:
    def _run_main_access_denied(self, monkeypatch):
        client, stubber = _stubbed_ce(monkeypatch)
        stubber.add_client_error(
            "get_cost_and_usage",
            service_error_code="AccessDeniedException",
            service_message="User is not authorized to perform: ce:GetCostAndUsage",
            http_status_code=400,
        )
        utils = billing_export.utils
        monkeypatch.setattr(utils, "detect_partition", lambda *a, **k: "aws")
        monkeypatch.setattr(utils, "print_script_banner", lambda *a, **k: ("123456789012", "TEST-ACCOUNT"))
        monkeypatch.setattr(utils, "ensure_dependencies", lambda *a, **k: True)
        monkeypatch.setattr(utils, "is_auto_run", lambda: True)
        warnings = []
        monkeypatch.setattr(utils, "log_warning", lambda msg: warnings.append(msg))

        with stubber, pytest.raises(SystemExit) as exc:
            billing_export.main()
        return exc.value.code, warnings

    def test_exit_zero_marker_and_warning(self, monkeypatch, patch_output_dir):
        code, warnings = self._run_main_access_denied(monkeypatch)
        assert code == 0
        assert any("lacks Cost Explorer permissions" in w for w in warnings)

        files = list(patch_output_dir.glob("*.xlsx"))
        assert len(files) == 1
        marker = files[0]
        assert "billing-skipped-no-permission-export-" in marker.name
        assert not fnmatch.fnmatch(marker.name, REAL_EXPORT_GLOB)

        wb = load_workbook(marker)
        assert "Summary" not in wb.sheetnames
        assert wb.sheetnames == ["Skipped"]
        values = dict(wb["Skipped"].iter_rows(values_only=True))
        assert values["Status"] == "SKIPPED"
        assert values["Error Code"] == "AccessDeniedException"
        assert values["Required Permission"] == "ce:GetCostAndUsage"

    def test_marker_lands_in_smart_scan_zip(self, monkeypatch, patch_output_dir):
        executor_mod = _import_smart_scan_package()
        ExecutionResult, ScriptExecutor = executor_mod.ExecutionResult, executor_mod.ScriptExecutor

        pre_run = {str(p) for p in patch_output_dir.glob("*.xlsx")}
        start_epoch = datetime.datetime.now().timestamp() - 1
        code, _ = self._run_main_access_denied(monkeypatch)
        assert code == 0

        executor = ScriptExecutor({"billing_export.py"})
        found = executor._find_output_file("billing_export", (pre_run, start_epoch))
        assert found is not None
        assert "billing-skipped-no-permission" in Path(found).name

        now = datetime.datetime.now()
        result = ExecutionResult(
            script="billing_export.py", success=True, start_time=now, end_time=now,
            duration_seconds=0.0, return_code=code, output_file=found,
        )
        smart_scan_cli = _load_smart_scan_module()
        zip_path = smart_scan_cli._zip_export_files([result], "TEST-ACCOUNT")
        assert zip_path is not None
        with zipfile.ZipFile(zip_path) as zf:
            names = zf.namelist()
        assert names == [Path(found).name]
        assert not any(fnmatch.fnmatch(n, REAL_EXPORT_GLOB) for n in names)


class TestGovCloudSkip:
    def _run_main_govcloud(self, monkeypatch, account_info):
        utils = billing_export.utils
        monkeypatch.setattr(utils, "detect_partition", lambda *a, **k: "aws-us-gov")
        monkeypatch.setattr(utils, "get_account_info", account_info)

        def _no_ce(*a, **k):
            raise AssertionError("Cost Explorer must not be called in GovCloud")

        monkeypatch.setattr(utils, "get_boto3_client", _no_ce)
        warnings = []
        monkeypatch.setattr(utils, "log_warning", lambda msg: warnings.append(msg))
        with pytest.raises(SystemExit) as exc:
            billing_export.main()
        return exc.value.code, warnings

    def _assert_marker(self, out_dir, expected_prefix, expected_account_id):
        files = list(out_dir.glob("*.xlsx"))
        assert len(files) == 1
        marker = files[0]
        assert marker.name.startswith(f"{expected_prefix}-billing-skipped-govcloud-export-")
        assert not fnmatch.fnmatch(marker.name, REAL_EXPORT_GLOB)
        wb = load_workbook(marker)
        assert wb.sheetnames == ["Skipped"]
        values = dict(wb["Skipped"].iter_rows(values_only=True))
        assert values["Status"] == "SKIPPED"
        assert "aws-us-gov" in values["Reason"]
        assert "commercial" in values["Reason"]
        assert values["Account ID"] == expected_account_id

    def test_exit_zero_and_marker(self, monkeypatch, patch_output_dir):
        code, warnings = self._run_main_govcloud(
            monkeypatch, lambda: ("123456789012", "GOV-TEST")
        )
        assert code == 0
        assert any("GovCloud" in w for w in warnings)
        self._assert_marker(patch_output_dir, "GOV-TEST", "123456789012")

    def test_account_lookup_failure_still_exits_zero(self, monkeypatch, patch_output_dir):
        def _boom():
            raise RuntimeError("sts unreachable")

        code, warnings = self._run_main_govcloud(monkeypatch, _boom)
        assert code == 0
        assert any("sts unreachable" in w for w in warnings)
        self._assert_marker(patch_output_dir, "UNKNOWN-ACCOUNT", "UNKNOWN")


class TestTwelveMonthReport:
    """End-to-end: 'last 12' request -> 12 month keys, sheets and Summary rows."""

    def test_last_12_report_has_exactly_12_months(self, monkeypatch, patch_output_dir):
        client, stubber = _stubbed_ce(monkeypatch)
        with _frozen_now(datetime.datetime(2026, 1, 1, 0, 1)):
            _, _, start, end = billing_export.validate_date_input("last 12")
        months = _months_in(start, end)

        results = []
        for key in months:
            y, m = int(key[:4]), int(key[5:])
            nxt = datetime.datetime(y + 1, 1, 1) if m == 12 else datetime.datetime(y, m + 1, 1)
            results.append(_result(f"{key}-01", nxt.strftime("%Y-%m-%d"),
                                   [_group("Amazon EC2", "1.00")]))
        # Split across two pages to exercise the loop on the real window.
        stubber.add_response(
            "get_cost_and_usage",
            {"ResultsByTime": results[:7], "NextPageToken": "p2"},
            _base_request("2025-01-01", "2026-01-01"),
        )
        stubber.add_response(
            "get_cost_and_usage",
            {"ResultsByTime": results[7:]},
            {**_base_request("2025-01-01", "2026-01-01"), "NextPageToken": "p2"},
        )
        with stubber:
            data = billing_export.get_billing_data(start, end)
            stubber.assert_no_pending_responses()
        assert sorted(data) == months and len(data) == 12

        path = billing_export.create_excel_report(
            data, "TEST-ACCOUNT", "last-12-months",
            account_id="123456789012", start_date=start, end_date=end,
        )
        wb = load_workbook(path)
        month_sheets = wb.sheetnames[3:]
        assert wb.sheetnames[:3] == ["Summary", "About", "Savings Plans"]
        assert len(month_sheets) == 12
        assert month_sheets[0] == "Jan 2025" and month_sheets[-1] == "Dec 2025"

        rows = list(wb["Summary"].iter_rows(values_only=True))
        assert rows[0] == ("Month", "Total Cost (USD)")
        assert len(rows) == 1 + 12 + 1
        assert rows[1][0] == "January 2025" and rows[12][0] == "December 2025"
        assert rows[-1][0] == "Total All Months"
        assert rows[-1][1] == pytest.approx(12.0)

        about = dict(list(wb["About"].iter_rows(values_only=True))[1:])
        assert about["Period Start (inclusive)"] == "2025-01-01"
        assert about["Period End (inclusive)"] == "2025-12-31"


# --- Savings Plans breakdown (Issue #300) ---------------------------------

def _rt_request(start, end, token=None):
    req = {
        "TimePeriod": {"Start": start, "End": end},
        "Granularity": "MONTHLY",
        "Metrics": [billing_export.COST_METRIC],
        "GroupBy": [{"Type": "DIMENSION", "Key": "RECORD_TYPE"}],
    }
    if token:
        req["NextPageToken"] = token
    return req


def _sp_status_block(ws):
    """(Field, Value) rows above the first blank row."""
    block = {}
    for row in ws.iter_rows(values_only=True):
        if row[0] is None:
            break
        block.setdefault(row[0], row[1])
    return block


def _sp_table(ws):
    rows = list(ws.iter_rows(values_only=True))
    start = next(i for i, r in enumerate(rows) if r[0] == "Month")
    return rows[start:]


class TestRecordTypeBreakdown:
    def test_multi_page_aggregation(self, monkeypatch):
        _, stubber = _stubbed_ce(monkeypatch)
        start, end = "2026-07-01", "2026-09-01"
        stubber.add_response(
            "get_cost_and_usage",
            {
                "ResultsByTime": [
                    _result("2026-07-01", "2026-08-01", [
                        _group("Usage", "40.00"),
                        _group("SavingsPlanCoveredUsage", "100.00"),
                        _group("SavingsPlanNegation", "-100.00"),
                    ]),
                    _result("2026-08-01", "2026-09-01", [
                        _group("SavingsPlanCoveredUsage", "60.00"),
                        _group("SavingsPlanRecurringFee", "30.00"),
                    ]),
                ],
                "NextPageToken": "p2",
            },
            _rt_request(start, end),
        )
        stubber.add_response(
            "get_cost_and_usage",
            {
                "ResultsByTime": [
                    _result("2026-08-01", "2026-09-01", [
                        _group("SavingsPlanCoveredUsage", "40.00"),
                        _group("SavingsPlanNegation", "-100.00"),
                        _group("SavingsPlanRecurringFee", "12.50"),
                        _group("SavingsPlanUpfrontFee", "5.00"),
                    ]),
                ],
            },
            _rt_request(start, end, "p2"),
        )
        with stubber:
            data = billing_export.get_record_type_breakdown(
                datetime.datetime(2026, 7, 1), datetime.datetime(2026, 9, 1)
            )
            stubber.assert_no_pending_responses()

        assert data["2026-07"] == {
            "Usage": pytest.approx(40.0),
            "SavingsPlanCoveredUsage": pytest.approx(100.0),
            "SavingsPlanNegation": pytest.approx(-100.0),
        }
        aug = data["2026-08"]
        assert aug["SavingsPlanCoveredUsage"] == pytest.approx(100.0)
        assert aug["SavingsPlanRecurringFee"] == pytest.approx(42.5)
        assert aug["SavingsPlanNegation"] == pytest.approx(-100.0)
        assert aug["SavingsPlanUpfrontFee"] == pytest.approx(5.0)

        sp = billing_export.summarize_savings_plans(data)
        assert sp["status"] == billing_export.SP_STATUS_FOUND
        assert sp["months"]["2026-07"]["SavingsPlanRecurringFee"] == 0.0
        assert sp["unrecognized"] == {}

    def test_empty_month_is_kept(self, monkeypatch):
        _, stubber = _stubbed_ce(monkeypatch)
        stubber.add_response(
            "get_cost_and_usage",
            {"ResultsByTime": [_result("2026-08-01", "2026-09-01", [])]},
            _rt_request("2026-08-01", "2026-09-01"),
        )
        with stubber:
            data = billing_export.get_record_type_breakdown(
                datetime.datetime(2026, 8, 1), datetime.datetime(2026, 9, 1)
            )
        assert data == {"2026-08": {}}


class TestSummarizeSavingsPlans:
    def test_no_sp_record_types_is_none_state(self):
        sp = billing_export.summarize_savings_plans({
            "2026-08": {"Usage": 10.0, "Tax": 1.0, "Credit": -2.0},
        })
        assert sp["status"] == billing_export.SP_STATUS_NONE

    def test_empty_breakdown_is_none_state(self):
        assert billing_export.summarize_savings_plans({})["status"] == billing_export.SP_STATUS_NONE

    def test_console_spelling_matches(self):
        sp = billing_export.summarize_savings_plans({
            "2026-08": {"Savings Plan Covered Usage": 5.0, "savings plan negation": -5.0},
        })
        assert sp["status"] == billing_export.SP_STATUS_FOUND
        row = sp["months"]["2026-08"]
        assert row["SavingsPlanCoveredUsage"] == 5.0
        assert row["SavingsPlanNegation"] == -5.0

    def test_unrecognized_sp_type_reported_not_netted(self):
        sp = billing_export.summarize_savings_plans({
            "2026-08": {"SavingsPlanSomethingNew": 7.0, "Usage": 1.0},
        })
        assert sp["status"] == billing_export.SP_STATUS_FOUND
        assert sp["unrecognized"] == {"SavingsPlanSomethingNew": 7.0}
        assert sum(sp["months"]["2026-08"].values()) == 0.0


class TestSavingsPlansSheet:
    DATA = {"2026-08": {"Amazon EC2": 50.0}}

    def _report(self, sp_result):
        path = billing_export.create_excel_report(
            self.DATA, "TEST-ACCOUNT", "08-2026",
            account_id="123456789012",
            start_date=datetime.datetime(2026, 8, 1),
            end_date=datetime.datetime(2026, 9, 1),
            sp_result=sp_result,
        )
        return load_workbook(path)["Savings Plans"]

    def test_records_table(self, patch_output_dir):
        sp = billing_export.summarize_savings_plans({
            "2026-07": {"Usage": 3.0},
            "2026-08": {
                "SavingsPlanCoveredUsage": 100.0,
                "SavingsPlanNegation": -100.0,
                "SavingsPlanRecurringFee": 42.5,
                "SavingsPlanUpfrontFee": 5.0,
            },
        })
        ws = self._report(sp)
        block = _sp_status_block(ws)
        assert block["Status"] == billing_export.SP_STATUS_FOUND
        assert block["Caveat"] == billing_export.SP_ABSENCE_CAVEAT
        table = _sp_table(ws)
        assert table[0] == (
            "Month", "Covered Usage (USD)", "Negation (USD)",
            "Recurring Fee (USD)", "Upfront Fee (USD)", "Net (USD)",
        )
        assert table[1] == ("July 2026", 0.0, 0.0, 0.0, 0.0, 0.0)
        assert table[2][0] == "August 2026"
        assert table[2][1:] == (
            pytest.approx(100.0), pytest.approx(-100.0), pytest.approx(42.5),
            pytest.approx(5.0), pytest.approx(47.5),
        )
        assert table[-1][0] == "Total"
        assert table[-1][-1] == pytest.approx(47.5)

    def test_no_records_state_is_explicit(self, patch_output_dir):
        ws = self._report(billing_export.summarize_savings_plans({"2026-08": {"Usage": 50.0}}))
        block = _sp_status_block(ws)
        assert block["Status"] == "NO RECORDS"
        assert block["Detail"] == (
            "No Savings Plans records in this account's Cost Explorer data for this period."
        )
        assert "another account" in block["Caveat"] and "reseller" in block["Caveat"]
        values = list(ws.iter_rows(values_only=True))
        assert not any(r[0] == "Month" for r in values), "no zeroed table in the empty state"

    def test_failed_state_distinct_from_none(self, patch_output_dir):
        err = billing_export.ClientError(
            {"Error": {"Code": "AccessDeniedException", "Message": "denied"}},
            "GetCostAndUsage",
        )
        ws = self._report(billing_export.savings_plans_failure(err))
        block = _sp_status_block(ws)
        assert block["Status"] == "LOOKUP FAILED"
        assert block["Status"] != billing_export.SP_STATUS_NONE
        assert "UNKNOWN" in block["Detail"] and "not the same as 'no records'" in block["Detail"]
        assert block["Error Code"] == "AccessDeniedException"
        assert block["Error Message"] == "denied"

    def test_none_result_renders_as_failed_not_empty(self, patch_output_dir):
        block = _sp_status_block(self._report(None))
        assert block["Status"] == billing_export.SP_STATUS_FAILED
        assert block["Error Code"] == "NotQueried"


class TestMainSavingsPlansFailure:
    """Breakdown query failure must not fail the main billing export."""

    def _run_main(self, monkeypatch, sp_error_code):
        _, stubber = _stubbed_ce(monkeypatch)
        with _frozen_now(datetime.datetime(2026, 9, 15, 12, 0)):
            _, _, start, end = billing_export.validate_date_input("last 12")
        s, e = start.strftime("%Y-%m-%d"), end.strftime("%Y-%m-%d")
        stubber.add_response(
            "get_cost_and_usage",
            {"ResultsByTime": [_result("2026-08-01", "2026-09-01", [_group("Amazon EC2", "9.00")])]},
            _base_request(s, e),
        )
        stubber.add_client_error(
            "get_cost_and_usage",
            service_error_code=sp_error_code,
            service_message="breakdown failed",
            http_status_code=400,
            expected_params=_rt_request(s, e),
        )
        utils = billing_export.utils
        monkeypatch.setattr(utils, "detect_partition", lambda *a, **k: "aws")
        monkeypatch.setattr(utils, "print_script_banner", lambda *a, **k: ("123456789012", "TEST-ACCOUNT"))
        monkeypatch.setattr(utils, "ensure_dependencies", lambda *a, **k: True)
        monkeypatch.setattr(utils, "is_auto_run", lambda: True)
        warnings = []
        monkeypatch.setattr(utils, "log_warning", lambda msg: warnings.append(msg))

        with _frozen_now(datetime.datetime(2026, 9, 15, 12, 0)), stubber:
            billing_export.main()
            stubber.assert_no_pending_responses()
        return warnings

    @pytest.mark.parametrize("code", ["AccessDeniedException", "DataUnavailableException"])
    def test_export_succeeds_and_sheet_records_failure(self, monkeypatch, patch_output_dir, code):
        warnings = self._run_main(monkeypatch, code)
        assert any("Savings Plans breakdown lookup failed" in w and code in w for w in warnings)

        files = list(patch_output_dir.glob("*.xlsx"))
        assert len(files) == 1
        assert fnmatch.fnmatch(files[0].name, REAL_EXPORT_GLOB)
        wb = load_workbook(files[0])
        assert wb.sheetnames[:3] == ["Summary", "About", "Savings Plans"]
        summary = list(wb["Summary"].iter_rows(values_only=True))
        assert summary[-1] == ("Total All Months", pytest.approx(9.0))
        block = _sp_status_block(wb["Savings Plans"])
        assert block["Status"] == billing_export.SP_STATUS_FAILED
        assert block["Error Code"] == code
