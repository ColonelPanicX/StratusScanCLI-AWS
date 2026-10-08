#!/usr/bin/env python3
"""
Tests for workspaces_export.py.

moto coverage of the WorkSpaces API (moto 5.1.x, checked 10.07.2026):
  implemented     : create_workspaces, describe_workspaces, describe_tags,
                    describe_workspace_directories, register_workspace_directory
  NOT implemented : describe_workspace_bundles,
                    describe_workspaces_connection_status
                    (both raise NotImplementedError)
moto's describe_workspaces also ignores pagination (it never returns a
NextToken), so pagination and the 25-ID batch limit are covered with fake
clients, and the failure paths with injected errors. moto cannot catch an
unread-page bug, so the fake-client tests are the real pagination proof.
"""

import datetime
import logging
import sys
from pathlib import Path

import boto3
import botocore.exceptions
import pytest
from moto import mock_aws
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).parent.parent.parent / "scripts"))
import workspaces_export  # noqa: E402
from workspaces_export import (  # noqa: E402
    CREATED_DATE_NOT_EXPOSED,
    _scan_region,
    build_summary,
    map_power_state,
)

REGION = "us-east-1"

# boto3 returns tz-aware datetimes (tzlocal in practice); use a non-UTC offset
# so the UTC conversion is actually exercised.
LAST_CONN = datetime.datetime(2026, 10, 1, 8, 30, 0, tzinfo=datetime.timezone(datetime.timedelta(hours=-4)))
CHECKED = datetime.datetime(2026, 10, 7, 12, 0, 0, tzinfo=datetime.timezone.utc)


@pytest.fixture(autouse=True)
def fake_aws_credentials(monkeypatch):
    monkeypatch.setenv("AWS_ACCESS_KEY_ID", "testing")
    monkeypatch.setenv("AWS_SECRET_ACCESS_KEY", "testing")
    monkeypatch.setenv("AWS_SECURITY_TOKEN", "testing")
    monkeypatch.setenv("AWS_SESSION_TOKEN", "testing")
    monkeypatch.setenv("AWS_DEFAULT_REGION", REGION)


def _client_error(code="Throttling", op="DescribeWorkspaces"):
    return botocore.exceptions.ClientError({"Error": {"Code": code, "Message": "x"}}, op)


# ---------------------------------------------------------------------------
# Fake client: records calls, paginates by returning several pages per call.
# ---------------------------------------------------------------------------

class _FakePaginator:
    def __init__(self, client, op):
        self._client, self._op = client, op

    def paginate(self, **kwargs):
        self._client.calls.append((self._op, kwargs))
        err = self._client.errors.get(self._op)
        if err:
            raise err
        return iter(self._client.pages(self._op, kwargs))


class _FakeClient:
    def __init__(self, workspaces=None, workspace_pages=None, connections=None,
                 bundles=None, directories=None, tags=None, errors=None):
        # workspace_pages lets a test force multiple pages of DescribeWorkspaces
        self.workspace_pages = workspace_pages or [workspaces or []]
        self.connections = connections or {}
        self.bundles = bundles or {}
        self.directories = directories or {}
        self.tags = tags or {}
        self.errors = errors or {}
        self.calls = []

    def get_paginator(self, op):
        return _FakePaginator(self, op)

    def pages(self, op, kwargs):
        if op == 'describe_workspaces':
            return [{'Workspaces': p} for p in self.workspace_pages]
        if op == 'describe_workspaces_connection_status':
            ids = kwargs['WorkspaceIds']
            half = max(1, len(ids) // 2)  # split into 2 pages to exercise paging
            return [
                {'WorkspacesConnectionStatus': [self.connections[i] for i in chunk if i in self.connections]}
                for chunk in (ids[:half], ids[half:])
            ]
        if op == 'describe_workspace_bundles':
            return [{'Bundles': [self.bundles[i] for i in kwargs['BundleIds'] if i in self.bundles]}]
        if op == 'describe_workspace_directories':
            return [{'Directories': [self.directories[i] for i in kwargs['DirectoryIds'] if i in self.directories]}]
        raise AssertionError(f"unexpected paginator {op}")

    def describe_tags(self, ResourceId):  # noqa: N803 - mirrors the boto3 kwarg name
        self.calls.append(('describe_tags', {'ResourceId': ResourceId}))
        err = self.errors.get('describe_tags')
        if err:
            raise err
        return {'TagList': self.tags.get(ResourceId, [])}


def _ws(i, **over):
    ws = {
        'WorkspaceId': f'ws-{i:09d}',
        'DirectoryId': 'd-1234567890',
        'UserName': f'user{i}',
        'IpAddress': f'10.0.0.{i % 250}',
        'State': 'AVAILABLE',
        'BundleId': 'wsb-bundle001',
        'SubnetId': 'subnet-0123456789abcdef0',
        'ComputerName': f'COMP{i}',
        'RootVolumeEncryptionEnabled': True,
        'UserVolumeEncryptionEnabled': False,
        'VolumeEncryptionKey': 'arn:aws:kms:us-east-1:111122223333:key/abc',
        'WorkspaceProperties': {
            'RunningMode': 'AUTO_STOP',
            'RunningModeAutoStopTimeoutInMinutes': 60,
            'RootVolumeSizeGib': 80,
            'UserVolumeSizeGib': 50,
            'ComputeTypeName': 'STANDARD',
            'Protocols': ['PCOIP'],
            'OperatingSystemName': 'WINDOWS_10',
        },
    }
    ws.update(over)
    return ws


def _full_client(workspaces, **kw):
    conns = {w['WorkspaceId']: {'WorkspaceId': w['WorkspaceId'], 'ConnectionState': 'CONNECTED',
                                'LastKnownUserConnectionTimestamp': LAST_CONN,
                                'ConnectionStateCheckTimestamp': CHECKED} for w in workspaces}
    return _FakeClient(
        workspaces=workspaces,
        connections=conns,
        bundles={'wsb-bundle001': {
            'BundleId': 'wsb-bundle001', 'Name': 'Standard with Windows 10', 'Owner': 'AMAZON',
            # Bundle timestamps must never surface as a WorkSpace created date.
            'CreationTime': 'BUNDLE-CREATED', 'LastUpdatedTime': 'BUNDLE-UPDATED'}},
        directories={'d-1234567890': {'DirectoryId': 'd-1234567890',
                                      'DirectoryName': 'corp.example.com',
                                      'DirectoryType': 'SIMPLE_AD'}},
        tags={workspaces[0]['WorkspaceId']: [{'Key': 'Env', 'Value': 'prod'}]} if workspaces else {},
        **kw,
    )


def _use(monkeypatch, client):
    monkeypatch.setattr(workspaces_export.utils, 'get_boto3_client', lambda *a, **k: client)


# ---------------------------------------------------------------------------
# moto-backed tests
# ---------------------------------------------------------------------------

class TestMoto:
    @mock_aws
    def test_empty_region_returns_no_rows_and_no_warnings(self):
        result = _scan_region(REGION)
        assert result == {'rows': [], 'warnings': []}

    @mock_aws
    def test_workspace_collected_and_unimplemented_enrichment_is_stated(self):
        ec2 = boto3.client("ec2", region_name=REGION)
        vpc = ec2.create_vpc(CidrBlock="10.0.0.0/16")["Vpc"]["VpcId"]
        subnets = [
            ec2.create_subnet(VpcId=vpc, CidrBlock=f"10.0.{i}.0/24",
                              AvailabilityZone=f"{REGION}{az}")["Subnet"]["SubnetId"]
            for i, az in enumerate("ab")
        ]
        directory_id = boto3.client("ds", region_name=REGION).create_directory(
            Name="corp.example.com", Password="Passw0rd!x", Size="Small",
            VpcSettings={"VpcId": vpc, "SubnetIds": subnets},
        )["DirectoryId"]
        ws = boto3.client("workspaces", region_name=REGION)
        ws.register_workspace_directory(DirectoryId=directory_id, EnableWorkDocs=False)
        ws.create_workspaces(Workspaces=[{
            "DirectoryId": directory_id, "UserName": "bob", "BundleId": "wsb-12345678",
            "Tags": [{"Key": "Env", "Value": "dev"}],
            "WorkspaceProperties": {"RunningMode": "AUTO_STOP"},
        }])

        result = _scan_region(REGION)

        assert len(result['rows']) == 1
        row = result['rows'][0]
        assert row['User Name'] == 'bob'
        assert row['State'] == 'AVAILABLE'
        assert row['Power State'] == 'Powered On'
        assert row['Running Mode'] == 'AUTO_STOP'
        assert row['Directory Name'] == 'corp.example.com'
        assert row['Tags'] == 'Env:dev'
        assert row['Region'] == REGION
        assert row['Created Date'] == CREATED_DATE_NOT_EXPOSED
        # moto lacks these two APIs; the row must say so rather than invent data
        assert row['Connection State'] == 'Unavailable'
        assert row['Bundle Name'] == 'Unavailable'
        details = ' '.join(w['Detail'] for w in result['warnings'])
        assert 'DescribeWorkspacesConnectionStatus failed' in details
        assert 'DescribeWorkspaceBundles failed' in details


# ---------------------------------------------------------------------------
# Row content
# ---------------------------------------------------------------------------

class TestRows:
    def test_row_columns_and_values(self, monkeypatch):
        client = _full_client([_ws(1)])
        _use(monkeypatch, client)

        result = _scan_region(REGION)

        assert result['warnings'] == []
        row = result['rows'][0]
        assert row['Computer Name'] == 'COMP1'
        assert row['WorkSpace ID'] == 'ws-000000001'
        assert row['State'] == 'AVAILABLE'
        assert row['Power State'] == 'Powered On'
        assert row['Running Mode'] == 'AUTO_STOP'
        assert row['Auto-Stop Timeout (min)'] == 60
        assert row['Connection State'] == 'CONNECTED'
        assert row['Last Known User Connection'] == '2026-10-01 12:30:00 UTC'
        assert row['Connection State Checked'] == '2026-10-07 12:00:00 UTC'
        assert row['Operating System'] == 'WINDOWS_10'
        assert row['Compute Type'] == 'STANDARD'
        assert row['Root Volume (GiB)'] == 80
        assert row['User Volume (GiB)'] == 50
        assert row['Protocols'] == 'PCOIP'
        assert row['Root Volume Encrypted'] is True
        assert row['User Volume Encrypted'] is False
        assert row['Bundle Name'] == 'Standard with Windows 10'
        assert row['Bundle Owner'] == 'AMAZON'
        assert row['Directory Name'] == 'corp.example.com'
        assert row['Directory Type'] == 'SIMPLE_AD'
        assert row['Tags'] == 'Env:prod'

    def test_created_date_is_stated_and_bundle_timestamps_not_passed_off(self, monkeypatch):
        _use(monkeypatch, _full_client([_ws(1)]))
        row = _scan_region(REGION)['rows'][0]
        assert row['Created Date'] == CREATED_DATE_NOT_EXPOSED
        flat = ' '.join(str(v) for v in row.values())
        assert 'BUNDLE-CREATED' not in flat and 'BUNDLE-UPDATED' not in flat
        assert not any('Updated' in k for k in row)

    def test_missing_optional_fields_read_na_not_crash(self, monkeypatch):
        bare = {'WorkspaceId': 'ws-bare00001', 'State': 'STOPPED'}
        client = _FakeClient(workspaces=[bare])
        _use(monkeypatch, client)
        row = _scan_region(REGION)['rows'][0]
        assert row['Operating System'] == 'N/A'
        assert row['Running Mode'] == 'N/A'
        assert row['Bundle Name'] == 'N/A'
        assert row['Directory Name'] == 'N/A'
        assert row['Power State'] == 'Powered Off'

    def test_workspace_missing_from_connection_response_is_stated(self, monkeypatch):
        client = _full_client([_ws(1)])
        client.connections = {}
        _use(monkeypatch, client)
        row = _scan_region(REGION)['rows'][0]
        assert row['Connection State'] == 'Not returned by API'

    def test_unknown_connection_state_passes_through(self, monkeypatch):
        client = _full_client([_ws(1)])
        client.connections['ws-000000001']['ConnectionState'] = 'UNKNOWN'
        _use(monkeypatch, client)
        assert _scan_region(REGION)['rows'][0]['Connection State'] == 'UNKNOWN'

    def test_error_state_fields_surface(self, monkeypatch):
        ws = _ws(1, State='ERROR', ErrorCode='ClientUnreachable', ErrorMessage='no route')
        _use(monkeypatch, _full_client([ws]))
        row = _scan_region(REGION)['rows'][0]
        assert row['State'] == 'ERROR'
        assert row['Error Code'] == 'ClientUnreachable'
        assert row['Error Message'] == 'no route'


class TestPowerStateMapping:
    @pytest.mark.parametrize("state, expected", [
        ('AVAILABLE', 'Powered On'),
        ('STARTING', 'Starting'),
        ('STOPPING', 'Stopping'),
        ('STOPPED', 'Powered Off'),
    ])
    def test_documented_states(self, state, expected):
        assert map_power_state(state) == expected

    @pytest.mark.parametrize("state", [
        'PENDING', 'IMPAIRED', 'UNHEALTHY', 'REBOOTING', 'REBUILDING', 'RESTORING',
        'MAINTENANCE', 'ADMIN_MAINTENANCE', 'TERMINATING', 'TERMINATED', 'SUSPENDED',
        'UPDATING', 'ERROR',
    ])
    def test_other_states_are_not_guessed(self, state):
        mapped = map_power_state(state)
        assert mapped.startswith('Not determinable')
        assert state in mapped
        assert 'Powered' not in mapped

    def test_missing_state(self):
        assert map_power_state(None) == 'N/A'


# ---------------------------------------------------------------------------
# Pagination and batching
# ---------------------------------------------------------------------------

class TestPagination:
    def test_all_describe_workspaces_pages_are_read(self, monkeypatch):
        page1 = [_ws(i) for i in range(1, 4)]
        page2 = [_ws(i) for i in range(4, 6)]
        client = _full_client(page1 + page2)
        client.workspace_pages = [page1, page2]
        _use(monkeypatch, client)

        rows = _scan_region(REGION)['rows']

        assert {r['WorkSpace ID'] for r in rows} == {f'ws-{i:09d}' for i in range(1, 6)}

    def test_connection_status_batches_at_25_ids_and_reads_every_page(self, monkeypatch):
        workspaces = [_ws(i) for i in range(1, 31)]  # 30 -> 25 + 5
        client = _full_client(workspaces)
        _use(monkeypatch, client)

        rows = _scan_region(REGION)['rows']

        conn_calls = [kw for op, kw in client.calls if op == 'describe_workspaces_connection_status']
        assert [len(kw['WorkspaceIds']) for kw in conn_calls] == [25, 5]
        assert len(rows) == 30
        assert all(r['Connection State'] == 'CONNECTED' for r in rows), \
            "a connection-status page was not read"

    def test_bundle_and_directory_lookups_are_deduplicated(self, monkeypatch):
        client = _full_client([_ws(i) for i in range(1, 6)])
        _use(monkeypatch, client)
        _scan_region(REGION)
        bundle_calls = [kw for op, kw in client.calls if op == 'describe_workspace_bundles']
        dir_calls = [kw for op, kw in client.calls if op == 'describe_workspace_directories']
        assert bundle_calls == [{'BundleIds': ['wsb-bundle001']}]
        assert dir_calls == [{'DirectoryIds': ['d-1234567890']}]

    def test_tags_fetched_once_per_workspace(self, monkeypatch):
        client = _full_client([_ws(i) for i in range(1, 4)])
        _use(monkeypatch, client)
        _scan_region(REGION)
        tag_calls = [kw['ResourceId'] for op, kw in client.calls if op == 'describe_tags']
        assert sorted(tag_calls) == ['ws-000000001', 'ws-000000002', 'ws-000000003']


# ---------------------------------------------------------------------------
# Failure states
# ---------------------------------------------------------------------------

class TestFailureStates:
    def test_describe_workspaces_failure_raises_not_empty(self, monkeypatch):
        client = _FakeClient(errors={'describe_workspaces': _client_error('AccessDeniedException')})
        _use(monkeypatch, client)
        with pytest.raises(botocore.exceptions.ClientError):
            _scan_region(REGION)

    def test_client_creation_failure_raises(self, monkeypatch):
        def boom(*a, **k):
            raise _client_error('Throttling')
        monkeypatch.setattr(workspaces_export.utils, 'get_boto3_client', boom)
        with pytest.raises(botocore.exceptions.ClientError):
            _scan_region(REGION)

    @pytest.mark.parametrize("op, cell, label", [
        ('describe_workspaces_connection_status', 'Connection State', 'DescribeWorkspacesConnectionStatus'),
        ('describe_workspace_bundles', 'Bundle Name', 'DescribeWorkspaceBundles'),
        ('describe_workspace_directories', 'Directory Name', 'DescribeWorkspaceDirectories'),
    ])
    def test_enrichment_failure_keeps_rows_and_is_stated(self, monkeypatch, op, cell, label):
        client = _full_client([_ws(1), _ws(2)], errors={op: _client_error('AccessDeniedException', label)})
        _use(monkeypatch, client)

        result = _scan_region(REGION)

        assert len(result['rows']) == 2, "enrichment failure must not drop WorkSpaces"
        assert all(r[cell] == 'Unavailable' for r in result['rows'])
        assert any(label in w['Detail'] for w in result['warnings'])

    def test_tag_failure_is_stated_per_workspace(self, monkeypatch):
        client = _full_client([_ws(1)], errors={'describe_tags': _client_error('ResourceNotFoundException', 'DescribeTags')})
        _use(monkeypatch, client)

        result = _scan_region(REGION)

        assert result['rows'][0]['Tags'] == 'Unavailable'
        assert any('DescribeTags failed for ws-000000001' in w['Detail'] for w in result['warnings'])

    def test_malformed_workspace_is_skipped_and_reported(self, monkeypatch):
        client = _full_client([_ws(1), _ws(2)])
        _use(monkeypatch, client)
        original = workspaces_export._build_workspace_row

        def flaky(ws, *a, **k):
            if ws['WorkspaceId'] == 'ws-000000002':
                raise KeyError('boom')
            return original(ws, *a, **k)

        monkeypatch.setattr(workspaces_export, '_build_workspace_row', flaky)

        result = _scan_region(REGION)

        assert [r['WorkSpace ID'] for r in result['rows']] == ['ws-000000001']
        assert any('Skipped malformed WorkSpace ws-000000002' in w['Detail'] for w in result['warnings'])

    def test_unavailable_in_partition_skips_without_calling_aws(self, monkeypatch):
        monkeypatch.setattr(workspaces_export.utils, 'is_service_available_in_partition', lambda *a, **k: False)

        def must_not_call(*a, **k):
            raise AssertionError("no client expected")

        monkeypatch.setattr(workspaces_export.utils, 'get_boto3_client', must_not_call)
        assert _scan_region(REGION) == {'rows': [], 'warnings': []}


# ---------------------------------------------------------------------------
# _run_export orchestration
# ---------------------------------------------------------------------------

class _Capture:
    def __init__(self):
        self.sheets = None
        self.filename = None
        self.failures = None


@pytest.fixture
def capture(monkeypatch):
    cap = _Capture()

    def fake_save(sheets, filename, prepare=False):
        cap.sheets, cap.filename = sheets, filename
        return f"/tmp/{filename}"

    monkeypatch.setattr(workspaces_export.utils, 'save_multiple_dataframes_to_excel', fake_save)

    def fake_report(account, rtype, failed):
        cap.failures = (account, rtype, failed)

    monkeypatch.setattr(workspaces_export.utils, 'report_collection_failures', fake_report)
    return cap


class TestRunExport:
    def test_all_regions_empty_exits_zero_without_file(self, monkeypatch, capture):
        _use(monkeypatch, _FakeClient())
        with pytest.raises(SystemExit) as exc:
            workspaces_export._run_export('111122223333', 'ACME', [REGION], None, 'all')
        assert exc.value.code == 0
        assert capture.sheets is None

    def test_writes_summary_workspaces_and_warnings_sheets(self, monkeypatch, capture):
        _use(monkeypatch, _full_client([_ws(1), _ws(2, State='STOPPED')]))

        workspaces_export._run_export('111122223333', 'ACME', [REGION], None, 'all')

        assert list(capture.sheets) == ['Summary', 'WorkSpaces', 'Collection Warnings']
        assert len(capture.sheets['WorkSpaces']) == 2
        assert 'workspaces' in capture.filename and 'all-us-east-1' in capture.filename
        assert capture.failures is None

    def test_state_filter_applies(self, monkeypatch, capture):
        _use(monkeypatch, _full_client([_ws(1), _ws(2, State='STOPPED')]))

        workspaces_export._run_export('111122223333', 'ACME', [REGION], 'STOPPED', 'stopped')

        df = capture.sheets['WorkSpaces']
        assert list(df['State']) == ['STOPPED']
        assert 'stopped' in capture.filename

    def test_failed_region_exits_nonzero_and_writes_marker(self, monkeypatch, capture):
        _use(monkeypatch, _FakeClient(errors={'describe_workspaces': _client_error('AccessDeniedException')}))

        with pytest.raises(SystemExit) as exc:
            workspaces_export._run_export('111122223333', 'ACME', [REGION], None, 'all')

        assert exc.value.code == 1
        assert capture.failures is not None
        assert [r for r, _ in capture.failures[2]] == [REGION]
        assert capture.sheets is None, "nothing succeeded, so no workbook"

    def test_partial_failure_still_exports_then_exits_nonzero(self, monkeypatch, capture):
        good, bad = 'us-east-1', 'us-west-2'

        def per_region(service, region_name=None, **k):
            if region_name == bad:
                return _FakeClient(errors={'describe_workspaces': _client_error('Throttling')})
            return _full_client([_ws(1)])

        monkeypatch.setattr(workspaces_export.utils, 'get_boto3_client', per_region)

        with pytest.raises(SystemExit) as exc:
            workspaces_export._run_export('111122223333', 'ACME', [good, bad], None, 'all')

        assert exc.value.code == 1
        assert len(capture.sheets['WorkSpaces']) == 1
        assert [r for r, _ in capture.failures[2]] == [bad]


# ---------------------------------------------------------------------------
# Summary
# ---------------------------------------------------------------------------

class TestSummary:
    def test_counts_by_state_os_mode_and_compute(self):
        import pandas as pd
        df = pd.DataFrame([
            {'State': 'AVAILABLE', 'Power State': 'Powered On', 'Operating System': 'WINDOWS_10',
             'Running Mode': 'AUTO_STOP', 'Compute Type': 'STANDARD', 'Connection State': 'CONNECTED', 'Region': REGION},
            {'State': 'AVAILABLE', 'Power State': 'Powered On', 'Operating System': 'UBUNTU_22_04',
             'Running Mode': 'ALWAYS_ON', 'Compute Type': 'STANDARD', 'Connection State': 'DISCONNECTED', 'Region': REGION},
            {'State': 'STOPPED', 'Power State': 'Powered Off', 'Operating System': 'WINDOWS_10',
             'Running Mode': 'AUTO_STOP', 'Compute Type': 'POWER', 'Connection State': 'UNKNOWN', 'Region': REGION},
        ])

        rows = build_summary(df, [REGION], [], [])
        by = {(r['Category'], r['Metric']): r['Value'] for r in rows}

        assert by[('Totals', 'Total WorkSpaces')] == 3
        assert by[('By State', 'AVAILABLE')] == 2
        assert by[('By State', 'STOPPED')] == 1
        assert by[('By Power State', 'Powered Off')] == 1
        assert by[('By Operating System', 'WINDOWS_10')] == 2
        assert by[('By Running Mode', 'ALWAYS_ON')] == 1
        assert by[('By Compute Type', 'POWER')] == 1
        assert by[('By Connection State', 'UNKNOWN')] == 1
        assert by[('Notes', 'Created Date')] == CREATED_DATE_NOT_EXPOSED

    def test_failed_regions_and_warnings_are_counted(self):
        import pandas as pd
        rows = build_summary(pd.DataFrame(), [REGION, 'us-west-2'], [('us-west-2', 'x')], [{'Region': REGION, 'Detail': 'd'}])
        by = {r['Metric']: r['Value'] for r in rows if r['Category'] == 'Totals'}
        assert by['Regions Failed'] == 1
        assert by['Collection Warnings'] == 1
        assert by['Total WorkSpaces'] == 0


# ---------------------------------------------------------------------------
# Regions where WorkSpaces is not offered
# ---------------------------------------------------------------------------

def _endpoint_error(region):
    return botocore.exceptions.EndpointConnectionError(
        endpoint_url=f"https://workspaces.{region}.amazonaws.com/"
    )


class TestUnsupportedRegions:
    UNSUPPORTED = "eu-north-1"  # not in the documented region list

    def test_unlisted_region_with_endpoint_error_is_stated_skip(self, monkeypatch):
        client = _FakeClient(errors={'describe_workspaces': _endpoint_error(self.UNSUPPORTED)})
        _use(monkeypatch, client)

        result = _scan_region(self.UNSUPPORTED)

        assert result['rows'] == []
        assert len(result['warnings']) == 1
        assert result['warnings'][0]['Region'] == self.UNSUPPORTED
        assert 'Not offered in region' in result['warnings'][0]['Detail']

    def test_listed_region_with_endpoint_error_still_fails(self, monkeypatch):
        client = _FakeClient(errors={'describe_workspaces': _endpoint_error(REGION)})
        _use(monkeypatch, client)
        with pytest.raises(botocore.exceptions.EndpointConnectionError):
            _scan_region(REGION)

    def test_unlisted_region_with_other_error_still_fails(self, monkeypatch):
        client = _FakeClient(errors={'describe_workspaces': _client_error('AccessDeniedException')})
        _use(monkeypatch, client)
        with pytest.raises(botocore.exceptions.ClientError):
            _scan_region(self.UNSUPPORTED)

    def test_unlisted_region_that_works_is_collected(self, monkeypatch):
        """A newly launched region missing from the docs list is not skipped."""
        _use(monkeypatch, _full_client([_ws(1)]))
        result = _scan_region(self.UNSUPPORTED)
        assert len(result['rows']) == 1
        assert result['warnings'] == []

    def test_documented_regions_include_known_gov_and_commercial(self):
        for r in ('us-east-1', 'us-gov-west-1', 'us-gov-east-1'):
            assert r in workspaces_export.DOCUMENTED_REGIONS

    def test_clean_run_with_only_unsupported_region_exits_zero(self, monkeypatch, capture):
        _use(monkeypatch, _FakeClient(errors={'describe_workspaces': _endpoint_error(self.UNSUPPORTED)}))
        with pytest.raises(SystemExit) as exc:
            workspaces_export._run_export('111122223333', 'ACME', [self.UNSUPPORTED], None, 'all')
        assert exc.value.code == 0
        assert capture.failures is None

    def test_mixed_run_supported_data_plus_unsupported_exits_zero_with_warning(self, monkeypatch, capture):
        def per_region(service, region_name=None, **k):
            if region_name == self.UNSUPPORTED:
                return _FakeClient(errors={'describe_workspaces': _endpoint_error(region_name)})
            return _full_client([_ws(1)])

        monkeypatch.setattr(workspaces_export.utils, 'get_boto3_client', per_region)

        workspaces_export._run_export('111122223333', 'ACME', [REGION, self.UNSUPPORTED], None, 'all')

        assert capture.failures is None
        assert len(capture.sheets['WorkSpaces']) == 1
        details = list(capture.sheets['Collection Warnings']['Detail'])
        assert any('Not offered in region' in d for d in details)

    def test_mixed_run_unsupported_skip_plus_real_failure_exits_one(self, monkeypatch, capture):
        failing = 'us-west-2'

        def per_region(service, region_name=None, **k):
            if region_name == self.UNSUPPORTED:
                return _FakeClient(errors={'describe_workspaces': _endpoint_error(region_name)})
            if region_name == failing:
                return _FakeClient(errors={'describe_workspaces': _client_error('AccessDeniedException')})
            return _full_client([_ws(1)])

        monkeypatch.setattr(workspaces_export.utils, 'get_boto3_client', per_region)

        with pytest.raises(SystemExit) as exc:
            workspaces_export._run_export('111122223333', 'ACME', [REGION, self.UNSUPPORTED, failing], None, 'all')

        assert exc.value.code == 1
        assert [r for r, _ in capture.failures[2]] == [failing], "unsupported region must not be marked failed"


# ---------------------------------------------------------------------------
# Issue #310: tz-aware datetimes mixed with N/A must survive a REAL save.
# Earlier tests used string timestamps and a faked save helper, so the real
# writer was never exercised.
# ---------------------------------------------------------------------------

@pytest.fixture
def real_out_dir(tmp_path, monkeypatch):
    d = tmp_path / "output"
    d.mkdir()
    monkeypatch.setattr(workspaces_export.utils, "logger", logging.getLogger("test-workspaces"))
    monkeypatch.setattr(workspaces_export.utils, "get_output_dir", lambda: d)
    monkeypatch.setattr(workspaces_export.utils, "deliver_output", lambda p: p)
    return d


class TestFmtTs:
    def test_aware_datetime_converted_to_utc_string(self):
        assert workspaces_export._fmt_ts(LAST_CONN) == '2026-10-01 12:30:00 UTC'

    def test_naive_datetime_formatted_as_is(self):
        assert workspaces_export._fmt_ts(datetime.datetime(2026, 1, 2, 3, 4, 5)) == '2026-01-02 03:04:05 UTC'

    @pytest.mark.parametrize("empty", [None, '', 0])
    def test_absent_is_na(self, empty):
        assert workspaces_export._fmt_ts(empty) == 'N/A'

    def test_no_row_value_is_a_datetime(self, monkeypatch):
        _use(monkeypatch, _full_client([_ws(1)]))
        row = _scan_region(REGION)['rows'][0]
        assert not any(isinstance(v, datetime.datetime) for v in row.values())


class TestRealWorkbook:
    def _mixed_client(self):
        """First WorkSpace never connected (N/A leads the column), later ones have datetimes."""
        workspaces = [_ws(1, State='STOPPED'), _ws(2), _ws(3)]
        client = _full_client(workspaces)
        client.connections.pop('ws-000000001')  # absent from the API -> not a datetime
        client.connections['ws-000000002'].pop('LastKnownUserConnectionTimestamp')  # never connected
        return client

    def test_mixed_na_and_aware_datetimes_write_and_reread(self, monkeypatch, real_out_dir):
        _use(monkeypatch, self._mixed_client())

        workspaces_export._run_export('111122223333', 'ACME', [REGION], None, 'all')

        files = list(real_out_dir.glob('*.xlsx'))
        assert len(files) == 1, "export must produce a workbook (Issue #310)"
        wb = load_workbook(files[0])
        assert wb.sheetnames == ['Summary', 'WorkSpaces', 'Collection Warnings']

        ws = wb['WorkSpaces']
        header = [c.value for c in ws[1]]
        last_i = header.index('Last Known User Connection')
        chk_i = header.index('Connection State Checked')
        id_i = header.index('WorkSpace ID')
        by_id = {r[id_i]: r for r in ws.iter_rows(min_row=2, values_only=True)}

        assert by_id['ws-000000001'][last_i] == 'N/A'
        assert by_id['ws-000000002'][last_i] == 'N/A'
        assert by_id['ws-000000003'][last_i] == '2026-10-01 12:30:00 UTC'
        assert by_id['ws-000000003'][chk_i] == '2026-10-07 12:00:00 UTC'

    def test_all_three_sheets_roundtrip_with_content(self, monkeypatch, real_out_dir):
        # Enrichment failure guarantees a non-empty Collection Warnings sheet.
        client = self._mixed_client()
        client.errors['describe_workspace_bundles'] = _client_error('AccessDeniedException', 'DescribeWorkspaceBundles')
        _use(monkeypatch, client)

        workspaces_export._run_export('111122223333', 'ACME', [REGION], None, 'all')

        wb = load_workbook(next(real_out_dir.glob('*.xlsx')))
        summary = {(r[0], r[1]): r[2] for r in wb['Summary'].iter_rows(min_row=2, values_only=True)}
        assert summary[('Totals', 'Total WorkSpaces')] == 3
        assert summary[('By State', 'STOPPED')] == 1
        assert wb['WorkSpaces'].max_row == 4  # header + 3
        warn_rows = list(wb['Collection Warnings'].iter_rows(min_row=2, values_only=True))
        assert any('DescribeWorkspaceBundles failed' in r[1] for r in warn_rows)

    def test_empty_warnings_sheet_still_writes(self, monkeypatch, real_out_dir):
        _use(monkeypatch, _full_client([_ws(1)]))
        workspaces_export._run_export('111122223333', 'ACME', [REGION], None, 'all')
        wb = load_workbook(next(real_out_dir.glob('*.xlsx')))
        assert 'Collection Warnings' in wb.sheetnames
